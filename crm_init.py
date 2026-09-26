#!/usr/bin/env python3
import csv, re, hashlib, datetime, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = pathlib.Path('/root/projeto-sites-demos')
PROSPECTS = ROOT / 'prospects.csv'
CRM_CSV = ROOT / 'projeto_sites_crm.csv'
ACT_CSV = ROOT / 'projeto_sites_atividades.csv'
XLSX = ROOT / 'projeto_sites_crm.xlsx'
NOW = datetime.datetime.now(datetime.timezone.utc).isoformat(timespec='seconds')

CRM_FIELDS = [
    'lead_id','created_at','updated_at','nome','telefone','wa_link','etapa','motivo_etapa',
    'origem','fonte_url','demo_url','whatsapp_status','regiao','vulnerabilidade','site_status',
    'abordagem_status','last_outbound_at','last_inbound_at','outbound_count','followup_count',
    'next_followup_at','cadence_level','last_message_id','needs_human','persona','notes','valor','closed_at'
]
ACT_FIELDS = ['ts','lead_id','nome','telefone','tipo','etapa_de','etapa_para','mensagem_id','resumo','proxima_acao']
STAGES = ['novo','enviado','qualificando','negociando','follow_up','converteu','nao_converteu','desativado']

def digits(s): return re.sub(r'\D','',s or '')
def lead_id(nome, telefone):
    base = (digits(telefone) or nome.lower()).encode('utf-8')
    return 'ps_' + hashlib.sha1(base).hexdigest()[:10]
def demo_from_evidence(e):
    m = re.search(r'https://hlreboques-oss\.github\.io/projeto-sites-demos/[^\s|,]+/?', e or '')
    return m.group(0) if m else ''
def msg_id(e):
    m = re.search(r'id=([A-Za-z0-9_-]+)', e or '')
    return m.group(1) if m else ''
def sent_time(e):
    m = re.search(r'Envio\s+([0-9T:\-+.]+)', e or '')
    return m.group(1) if m else ''
def stage_from_status(st):
    st = (st or '').lower()
    if 'falha' in st or 'nao_existe' in st or 'não_existe' in st: return 'desativado','número inválido/sem WhatsApp na rota verificada'
    if 'enviada' in st or 'enviado' in st: return 'enviado','primeira abordagem enviada'
    return 'novo','aguardando abordagem'

rows=[]
if PROSPECTS.exists():
    with PROSPECTS.open(newline='', encoding='utf-8') as f:
        for p in csv.DictReader(f):
            tel = p.get('telefone') or ''
            etapa, motivo = stage_from_status(p.get('abordagem_status'))
            ev = p.get('evidencia_textual') or ''
            demo = demo_from_evidence(ev)
            mid = msg_id(ev)
            sent = sent_time(ev)
            rows.append({
                'lead_id': lead_id(p.get('nome',''), tel),
                'created_at': sent or NOW,
                'updated_at': NOW,
                'nome': p.get('nome',''),
                'telefone': tel,
                'wa_link': p.get('wa_link',''),
                'etapa': etapa,
                'motivo_etapa': motivo,
                'origem': 'prospects.csv/importado',
                'fonte_url': p.get('fonte_url',''),
                'demo_url': demo,
                'whatsapp_status': p.get('whatsapp_status',''),
                'regiao': p.get('regiao',''),
                'vulnerabilidade': p.get('vulnerabilidade',''),
                'site_status': p.get('site_status',''),
                'abordagem_status': p.get('abordagem_status',''),
                'last_outbound_at': sent,
                'last_inbound_at': '',
                'outbound_count': '1' if sent and etapa!='desativado' else '0',
                'followup_count': '0',
                'next_followup_at': '',
                'cadence_level': '0',
                'last_message_id': mid,
                'needs_human': 'false',
                'persona': 'Mayara',
                'notes': '',
                'valor': '',
                'closed_at': ''
            })

# write csvs
with CRM_CSV.open('w', newline='', encoding='utf-8') as f:
    w=csv.DictWriter(f, fieldnames=CRM_FIELDS); w.writeheader(); w.writerows(rows)
if not ACT_CSV.exists():
    with ACT_CSV.open('w', newline='', encoding='utf-8') as f:
        w=csv.DictWriter(f, fieldnames=ACT_FIELDS); w.writeheader()
        for r in rows:
            if r['etapa']=='enviado':
                w.writerow({'ts':r['last_outbound_at'] or r['created_at'],'lead_id':r['lead_id'],'nome':r['nome'],'telefone':r['telefone'],'tipo':'outbound_inicial','etapa_de':'novo','etapa_para':'enviado','mensagem_id':r['last_message_id'],'resumo':'Primeira abordagem enviada antes da criação do CRM.','proxima_acao':'monitorar resposta'})
            elif r['etapa']=='desativado':
                w.writerow({'ts':r['updated_at'],'lead_id':r['lead_id'],'nome':r['nome'],'telefone':r['telefone'],'tipo':'sistema','etapa_de':'novo','etapa_para':'desativado','mensagem_id':'','resumo':r['motivo_etapa'],'proxima_acao':'não contatar'})

# workbook
wb=Workbook(); ws=wb.active; ws.title='CRM'
ws.append(CRM_FIELDS)
for r in rows: ws.append([r.get(c,'') for c in CRM_FIELDS])

colors = {'novo':'D9EAF7','enviado':'FFF2CC','qualificando':'D9EAD3','negociando':'D9D2E9','follow_up':'FCE5CD','converteu':'B6D7A8','nao_converteu':'EADCF8','desativado':'E6B8AF'}
header_fill=PatternFill('solid', fgColor='1F4E78')
header_font=Font(color='FFFFFF', bold=True)
for c in ws[1]:
    c.fill=header_fill; c.font=header_font; c.alignment=Alignment(horizontal='center')
for row in ws.iter_rows(min_row=2):
    etapa=row[6].value
    fill=PatternFill('solid', fgColor=colors.get(etapa,'FFFFFF'))
    for c in row: c.fill=fill; c.alignment=Alignment(vertical='top', wrap_text=True)
for idx, width in enumerate([15,21,21,32,17,30,15,28,22,42,48,18,32,44,44,28,21,21,14,14,21,12,24,12,12,44,12,21], start=1):
    ws.column_dimensions[get_column_letter(idx)].width=width
ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
if rows:
    tab=Table(displayName='TabelaCRM', ref=f'A1:{get_column_letter(len(CRM_FIELDS))}{len(rows)+1}')
    tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tab)

kan=wb.create_sheet('Kanban')
kan.append(['Etapa','Quantidade','Critério operacional'])
criteria={
'novo':'lead encontrado, ainda sem envio',
'enviado':'primeira abordagem enviada e verificada',
'qualificando':'respondeu; entender contexto/necessidade',
'negociando':'conversa comercial ativa; preço/call/proposta',
'follow_up':'não comprou ainda ou esfriou; cadência espaçada até 15 dias',
'converteu':'pagou/fechou',
'nao_converteu':'disse não ou perdeu interesse sem opt-out definitivo',
'desativado':'não contatar: opt-out, empresa fechada, número inválido, já tem site e não quer IA/site'
}
for st in STAGES:
    kan.append([st, f'=COUNTIF(CRM!G:G,"{st}")', criteria[st]])
for c in kan[1]: c.fill=header_fill; c.font=header_font
kan.column_dimensions['A'].width=18; kan.column_dimensions['B'].width=12; kan.column_dimensions['C'].width=90

act=wb.create_sheet('Atividades')
with ACT_CSV.open(newline='', encoding='utf-8') as f:
    for row in csv.reader(f): act.append(row)
for c in act[1]: c.fill=header_fill; c.font=header_font
for i,w in enumerate([21,15,32,17,18,16,16,24,60,40], start=1): act.column_dimensions[get_column_letter(i)].width=w
act.freeze_panes='A2'

cad=wb.create_sheet('Cadencia')
cad_rows=[
 ['Nível','Quando usar','Espaçamento mínimo','Mensagem/ação'],
 ['0','Primeira abordagem enviada','aguardar resposta','não reenviar só para bater volume'],
 ['1','não respondeu após 2 dias','2 dias','lembrar que a prévia é gratuita/não oficial e perguntar se quer que personalize com fotos reais'],
 ['2','sem resposta após follow-up 1','4 dias','mostrar benefício específico do nicho; mensagem curta'],
 ['3','sem resposta após follow-up 2','7 dias','última tentativa educada; perguntar se faz sentido retomar outro momento'],
 ['4','sem resposta após 15 dias','15 dias','encerrar cadência e marcar nao_converteu ou frio']
]
for row in cad_rows: cad.append(row)
for c in cad[1]: c.fill=header_fill; c.font=header_font
for i,w in enumerate([10,36,20,100], start=1): cad.column_dimensions[get_column_letter(i)].width=w

reg=wb.create_sheet('Regras')
reg_rows=[
 ['Regra','Descrição'],
 ['Persona','Mayara: agente de IA que pesquisa clientes na internet, identifica oportunidades e cria prévias personalizadas.'],
 ['Site existente','Se disser que já tem site, oferecer solução de IA/automação/atendimento/auditoria; se não interessar, desativar.'],
 ['Empresa parada/fechada','Tirar da jogada: etapa desativado.'],
 ['Não interessado/opt-out','Parar contato: etapa desativado.'],
 ['Não comprou mas abriu conversa','Mover para follow_up com cadência espaçada, sem ser chato, no máximo até 15 dias.'],
 ['Comprou','Mover para converteu e registrar valor/observação.'],
 ['Não comprou definitivo','Mover para nao_converteu.'],
 ['Personalização','Demos devem evitar genérico e deixar claro que fotos reais do estabelecimento podem ser incluídas com aprovação.']
]
for row in reg_rows: reg.append(row)
for c in reg[1]: c.fill=header_fill; c.font=header_font
reg.column_dimensions['A'].width=24; reg.column_dimensions['B'].width=120

wb.save(XLSX)
print({'crm_csv':str(CRM_CSV),'atividades_csv':str(ACT_CSV),'xlsx':str(XLSX),'rows':len(rows)})
